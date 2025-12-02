import os
from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QScrollArea, QLabel, QLineEdit,
    QCheckBox, QButtonGroup, QPushButton, QFrame, QSizePolicy, QCompleter
)
from PySide6.QtCore import Qt, QEvent, QTimer
from PySide6.QtGui import QFontMetrics
import openpyxl
from utilities import require_file

EXCEL_PATH = r"P:\ENGINEERING\Design Checklist\supporting_documents\checklist_questions.xlsx"


# Cap for the four top text fields (adjust if you want them a bit wider/narrower)
TOP_FIELD_MAX_W = 380

# ---- Auto width with dynamic max (passed in at call time) ----
def auto_resize_lineedit(lineedit, min_width=120, extra=18, max_width=None):
    text = lineedit.text() or lineedit.placeholderText() or ""
    fm = QFontMetrics(lineedit.font())
    new_width = fm.horizontalAdvance(text) + extra
    if max_width is None:
        new_width = max(min_width, new_width)
    else:
        new_width = max(min_width, min(new_width, max_width))
    if abs(lineedit.width() - new_width) > 2:
        lineedit.setFixedWidth(new_width)

class ChecklistTab(QWidget):
    def __init__(self, dirty_tracker=None, parent=None):
        super().__init__(parent)
        self.dirty_tracker = dirty_tracker
        self._loading = False
        self.category_frames = {}
        self.question_widgets = {}
        self.top_fields = []
        self.init_ui()

    # ----- helper: how wide can the top fields be right now? -----
    def _available_topfield_width(self):
        # Keep a practical cap so they never stretch across the whole tab
        return min(TOP_FIELD_MAX_W, max(180, self.width() - 40))

    def init_ui(self):
        main_layout = QVBoxLayout(self)
        main_layout.setContentsMargins(4, 4, 4, 4)
        main_layout.setSpacing(6)

        header = QWidget()
        header_layout = QVBoxLayout(header)
        header_layout.setContentsMargins(0, 0, 0, 0)
        header_layout.setSpacing(3)

        # ----- Top fields (UI order: ID, Customer, Opp, Sales) -----
        top_fields_row = QWidget()
        top_fields_layout = QVBoxLayout(top_fields_row)
        top_fields_layout.setContentsMargins(0, 0, 0, 0)
        top_fields_layout.setSpacing(8)
        placeholders = ["ID #", "Customer Name", "Opp Name", "Sales/CSR"]
        self.top_fields = []
        for text in placeholders:
            field = QLineEdit()
            field.setPlaceholderText(text)
            field.setStyleSheet("font-size: 13px;")
            # Do NOT let them expand; we’ll drive width ourselves
            field.setSizePolicy(QSizePolicy.Fixed, QSizePolicy.Fixed)
            field.setFixedWidth(140)  # initial; will be autosized below
            field.textChanged.connect(lambda _, f=field: self.on_topfield_changed(f))
            self.top_fields.append(field)
            top_fields_layout.addWidget(field)
        header_layout.addWidget(top_fields_row)
        main_layout.addWidget(header)

        # Initial cap & sizing for top fields
        self.update_topfield_caps()

        # ----- Sales/CSR autocomplete -----
        sales_file = r"P:\ENGINEERING\Design Checklist\supporting_documents\sales_list.txt"
        if os.path.exists(sales_file):
            with open(sales_file, "r", encoding="utf-8") as f:
                names = [line.strip() for line in f if line.strip()]
            names.sort(key=lambda s: s.lower())

            comp = QCompleter(names)
            comp.setCaseSensitivity(Qt.CaseInsensitive)
            comp.setFilterMode(Qt.MatchContains)
            comp.setCompletionMode(QCompleter.PopupCompletion)
            comp.setMaxVisibleItems(15)

            sales_field = self.top_fields[3]  # Sales/CSR
            sales_field.setCompleter(comp)
            sales_field.installEventFilter(self)

        # ----- Scrollable checklist area -----
        scroll_area = QScrollArea()
        scroll_area.setWidgetResizable(True)
        scroll_content = QWidget()
        scroll_layout = QVBoxLayout(scroll_content)
        scroll_layout.setContentsMargins(4, 4, 4, 4)
        scroll_layout.setSpacing(8)
        scroll_area.setWidget(scroll_content)
        main_layout.addWidget(scroll_area, stretch=1)

        # Load questions from Excel and group
        questions = self.load_questions_from_excel()
        grouped = {}
        for q in questions:
            grouped.setdefault(q["category"], []).append(q)

        # ----- General Questions -----
        general_frame = QFrame()
        general_layout = QVBoxLayout(general_frame)
        general_layout.setContentsMargins(4, 6, 4, 6)
        general_layout.setSpacing(5)
        general_label = QLabel("General Questions")
        general_label.setStyleSheet("font-weight: bold; font-size: 15px; text-decoration: underline; margin-bottom: 4px;")
        general_layout.addWidget(general_label)
        for q in grouped.get("General", []):
            self.add_question(general_layout, "General", q["question"], q["note"])
        scroll_layout.addWidget(general_frame)

        # ----- CD, MT, MIS side-by-side -----
        category_row = QWidget()
        category_layout = QHBoxLayout(category_row)
        category_layout.setContentsMargins(0, 0, 0, 0)
        category_layout.setSpacing(10)

        for category in ["CD", "MT", "MIS"]:
            frame = QFrame()
            frame.setLayout(QVBoxLayout())
            frame.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Preferred)
            frame.layout().setContentsMargins(4, 6, 4, 6)
            frame.layout().setSpacing(5)
            self.category_frames[category] = frame

            if category in grouped:
                header_lbl = QLabel(f"{category} Questions")
                header_lbl.setAlignment(Qt.AlignHCenter)
                header_lbl.setStyleSheet("font-weight: bold; text-decoration: underline; font-size: 14px; margin-bottom: 2px;")
                frame.layout().addWidget(header_lbl)

                for q in grouped[category]:
                    self.add_question(frame.layout(), category, q["question"], q["note"])

            category_layout.addWidget(frame, alignment=Qt.AlignTop)

        scroll_layout.addWidget(category_row)

    # Recompute caps & resize widths when the window resizes
    def resizeEvent(self, event):
        super().resizeEvent(event)
        self.update_topfield_caps()

    def showEvent(self, event):
        super().showEvent(event)
        # After first layout pass, ensure Qt doesn't stretch them
        QTimer.singleShot(0, self.update_topfield_caps)

    def update_topfield_caps(self):
        maxw = self._available_topfield_width()
        for f in self.top_fields:
            f.setMaximumWidth(maxw)
            f.setMinimumWidth(min(240, maxw))
            auto_resize_lineedit(f, min_width=140, extra=22, max_width=maxw)

    def on_topfield_changed(self, field):
        if self._loading:
            return
        maxw = self._available_topfield_width()
        auto_resize_lineedit(field, min_width=140, extra=22, max_width=maxw)
        if self.dirty_tracker:
            self.dirty_tracker.mark_dirty()

    def load_questions_from_excel(self):
        # If the Excel file with questions is missing, just return an empty list.
        # The tab will still load; you just won't see any questions.
        if not os.path.exists(EXCEL_PATH):
            return []

        wb = openpyxl.load_workbook(EXCEL_PATH)
        ws = wb.active
        questions = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            category, question, note = row[:3]
            if not category or not question:
                continue
            questions.append({
                "category": category.strip(),
                "question": question.strip(),
                "note": (note or "").strip()
            })
        return questions



    def add_question(self, parent_layout, category, question, note=None):
        wrapper = QWidget()
        layout = QVBoxLayout(wrapper)
        layout.setContentsMargins(2, 0, 2, 10)
        layout.setSpacing(2)

        label = QLabel(question)
        label.setStyleSheet("font-size: 13px; margin-bottom: 0;")
        layout.addWidget(label, alignment=Qt.AlignLeft)
        if note:
            note_lbl = QLabel(note)
            note_lbl.setStyleSheet("font-style: italic; color: gray; font-size: 11px; margin-bottom: 0;")
            layout.addWidget(note_lbl, alignment=Qt.AlignLeft)

        btn_group = QButtonGroup(wrapper)
        btn_group.setExclusive(False)
        btn_row = QWidget()
        btn_layout = QHBoxLayout(btn_row)
        btn_layout.setContentsMargins(0, 0, 0, 0)
        btn_layout.setSpacing(10)

        answers = ["Yes", "No", "N/A"]
        buttons = {}

        def on_btn_clicked(btn, qkey):
            if self._loading:
                return
            if btn.isChecked():
                for other in btn_group.buttons():
                    if other is not btn:
                        other.setChecked(False)
                self.question_widgets[qkey]["selected"] = btn.text()
            else:
                self.question_widgets[qkey]["selected"] = None
            if self.dirty_tracker:
                self.dirty_tracker.mark_dirty()

        qkey = (category, question)
        self.question_widgets[qkey] = {"selected": None, "buttons": {}, "btn_group": btn_group}

        for answer in answers:
            btn = QPushButton(answer)
            btn.setCheckable(True)
            btn.setStyleSheet("font-size: 13px; padding: 1px 12px;")
            btn.setFixedWidth(62)
            btn_group.addButton(btn)
            btn.clicked.connect(lambda checked, b=btn, qkey=qkey: on_btn_clicked(b, qkey))
            btn_layout.addWidget(btn)
            buttons[answer] = btn

        btn_layout.addStretch()
        layout.addWidget(btn_row, alignment=Qt.AlignLeft)

        self.question_widgets[qkey]["buttons"] = buttons
        parent_layout.addWidget(wrapper)

    # ---------- SAVE (keep legacy order for Saved Checklists) ----------
    def get_checklist_data(self):
        # UI order: [ID, Customer, Opp, Sales]
        id_txt    = self.top_fields[0].text().strip()
        cust_txt  = self.top_fields[1].text().strip()
        opp_txt   = self.top_fields[2].text().strip()
        sales_txt = self.top_fields[3].text().strip()

        # Legacy storage order expected elsewhere: [Customer, Opp, ID, Sales]
        top_fields_legacy = [cust_txt, opp_txt, id_txt, sales_txt]

        return {
            "top_fields": top_fields_legacy,
            "categories": {},
            "answers": {
                f"{cat}::{question}": data["selected"]
                for (cat, question), data in self.question_widgets.items()
                if data["selected"]
            }
        }

    # ---------- LOAD (legacy -> new UI order) ----------
    def load_checklist_data(self, data, read_only=False):
        self._loading = True
        self.setUpdatesEnabled(False)

        # Legacy order: [Customer, Opp, ID, Sales]
        tf = data.get("top_fields", [])
        cust_txt  = tf[0] if len(tf) > 0 else ""
        opp_txt   = tf[1] if len(tf) > 1 else ""
        id_txt    = tf[2] if len(tf) > 2 else ""
        sales_txt = tf[3] if len(tf) > 3 else ""

        # UI order -> [ID, Customer, Opp, Sales]
        values_in_ui_order = [id_txt, cust_txt, opp_txt, sales_txt]
        for i, value in enumerate(values_in_ui_order):
            if i < len(self.top_fields):
                field = self.top_fields[i]
                if field.text() != value:
                    field.setText(value)
                # Apply current cap during load
                maxw = self._available_topfield_width()
                auto_resize_lineedit(field, min_width=140, extra=22, max_width=maxw)
                field.setReadOnly(read_only)

        # Answers
        for key, answer in data.get("answers", {}).items():
            cat, question = key.split("::", 1)
            qkey = (cat, question)
            if qkey in self.question_widgets:
                self.question_widgets[qkey]["selected"] = answer
                btns = self.question_widgets[qkey]["buttons"]
                for btn in btns.values():
                    should_be_checked = btn.text() == answer
                    if btn.isChecked() != should_be_checked:
                        btn.setChecked(should_be_checked)
                    btn.setEnabled(not read_only)

        self.setUpdatesEnabled(True)
        self._loading = False

    def clear_checklist_tab(self):
        self._loading = True
        self.setUpdatesEnabled(False)

        for field in self.top_fields:
            if field.text():
                field.clear()
            if field.isReadOnly():
                field.setReadOnly(False)

        for qkey in self.question_widgets:
            self.question_widgets[qkey]["selected"] = None
            btns = self.question_widgets[qkey]["buttons"]
            for btn in btns.values():
                if btn.isChecked():
                    btn.setChecked(False)
                if not btn.isEnabled():
                    btn.setEnabled(True)

        self.setUpdatesEnabled(True)
        self._loading = False

    def eventFilter(self, obj, event):
        # When Sales/CSR gains focus, show all entries
        if obj is self.top_fields[3] and event.type() == QEvent.FocusIn:
            comp = obj.completer()
            if comp:
                comp.complete()
        return super().eventFilter(obj, event)


# ---- Factory functions for external API ----
_tab_instance = None

def create_checklist_tab(tab_widget: QWidget, dirty_tracker=None):
    global _tab_instance
    _tab_instance = ChecklistTab(dirty_tracker)
    layout = QVBoxLayout(tab_widget)
    layout.setContentsMargins(0, 0, 0, 0)
    layout.addWidget(_tab_instance)

def get_checklist_data():
    return _tab_instance.get_checklist_data() if _tab_instance else {}

def load_checklist_data(data, read_only=False):
    if _tab_instance:
        _tab_instance.load_checklist_data(data, read_only=read_only)

def clear_checklist_tab():
    if _tab_instance:
        _tab_instance.clear_checklist_tab()
